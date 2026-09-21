# CELL 0
from google.colab import drive
drive.mount('/content/drive/', force_remount=True)

import os
os.chdir('/content/drive/MyDrive/ve/variance_analysis')

# CELL 2
!pip install -e /content/drive/MyDrive/ve/variance_analysis

!pip -q install fonttools

!ttx -l fonts/Helvetica.ttc

!sudo mkdir -p /usr/share/fonts/truetype/helvetica
!ttx -y 0 -o /usr/share/fonts/truetype/helvetica/Helvetica-Regular.ttf fonts/Helvetica.ttc
!ttx -y 1 -o /usr/share/fonts/truetype/helvetica/Helvetica-Oblique.ttf fonts/Helvetica.ttc
!ttx -y 2 -o /usr/share/fonts/truetype/helvetica/Helvetica-Bold.ttf fonts/Helvetica.ttc
!ttx -y 3 -o /usr/share/fonts/truetype/helvetica/Helvetica-BoldOblique.ttf fonts/Helvetica.ttc

!fc-cache -fv


# CELL 3
import numpy as np
import matplotlib.pyplot as plt
import pandas as pd
from scipy.io import loadmat
from scipy.io.matlab._mio5_params import mat_struct
import os
from pathlib import Path
import matplotlib.font_manager as fm
import shutil

current_file_path = Path().resolve()
current_file_path

# set up matplotlib formatting too

!mkdir -p ~/.fonts
!cp fonts/Helvetica.ttc ~/.fonts/

cache_dir = os.path.join(os.path.expanduser("~"), ".cache", "matplotlib")
if os.path.exists(cache_dir):

    shutil.rmtree(cache_dir)

fm._load_fontmanager(try_read_cache=False)
plt.rcParams['font.family'] = 'Helvetica'


# CELL 4
fig, ax = plt.subplots(figsize=(4, 2), dpi=150)
ax.text(0.5, 0.62, 'The quick brown fox', ha='center', va='center', fontsize=16)
ax.axis('off')
plt.show()


# CELL 6

save_path = current_file_path / f'results/empirical/dataset_2'
data_path = current_file_path / f'data/empirical/dataset_2'

rats = [1,2,3,4]
subject_units = [204, 174, 79, 94]



# CELL 7
from src.dataset_2_decode_solo import decode_solo

reload = True

for subject in rats:
    try:
      if not reload and f'errors_v2_cv0_360_all.npz' in os.listdir(save_path / f'Subject{subject}'):
        continue
    except Exception as e:
      pass

    decode_solo(
        data_path=data_path,
        save_path=save_path,
        subject=subject,
        cv_folds=4,
        save_intermediates=True
      )


# CELL 8
use_median = True  # Use median for summary stats
ARENA_SIZE_CM = 200.0
V1_GRID = 100
V2_BINS = 360
bins = 60
density = True

# No-prior chance values (uniform random guessing)
import math
sqrt2 = math.sqrt(2)
NO_PRIOR_V1_COEFF = (2 + sqrt2 + 5 * math.log(1 + sqrt2)) / 15  # ~0.5214
NO_PRIOR_V1_CM = NO_PRIOR_V1_COEFF * ARENA_SIZE_CM  # ~104.28 cm for 200cm arena
NO_PRIOR_V2_DEG = 90.0  # Expected error for uniform guessing over [0, 360)

# ============ Imports ============
import numpy as np
import matplotlib.pyplot as plt
from pathlib import Path
from scipy.stats import gaussian_kde

# ============ Load data ============
v1_all = []
v2_all = []
v1_chance_all = []
v2_chance_all = []
subject_v1_native = []
subject_v2_native = []
subject_v1_chance_native = []
subject_v2_chance_native = []

for subject in rats:
    dpath = save_path / f"Subject{subject}"
    if not dpath.exists():
        continue

    # Neural decoding files
    xfiles = sorted(dpath.glob("errors_v1_cv*_*x*.npz"))
    hfiles = sorted(dpath.glob("errors_v2_cv*_360.npz"))

    # Chance decoding files
    xfiles_chance = sorted(dpath.glob("errors_v1_chance_cv*_*x*.npz"))
    hfiles_chance = sorted(dpath.glob("errors_v2_chance_cv*_360.npz"))

    # Load V1 neural decoding errors
    if xfiles:
        a = np.concatenate([np.asarray(np.load(f, allow_pickle=True)["errors_cm"]).ravel() for f in xfiles])
        a = a[np.isfinite(a)]
        if a.size:
            v1_all.append(a)
            subject_v1_native.append(np.nanmedian(a) if use_median else np.nanmean(a))

    # Load V1 chance decoding errors
    if xfiles_chance:
        a = np.concatenate([np.asarray(np.load(f, allow_pickle=True)["errors_cm"]).ravel() for f in xfiles_chance])
        a = a[np.isfinite(a)]
        if a.size:
            v1_chance_all.append(a)
            subject_v1_chance_native.append(np.nanmedian(a) if use_median else np.nanmean(a))

    # Load V2 neural decoding errors
    if hfiles:
        b = np.concatenate([np.abs(np.asarray(np.load(f, allow_pickle=True)["errors_deg"]).ravel()) for f in hfiles])
        b = b[np.isfinite(b)]
        if b.size:
            v2_all.append(b)
            subject_v2_native.append(np.nanmedian(b) if use_median else np.nanmean(b))

    # Load V2 chance decoding errors
    if hfiles_chance:
        b = np.concatenate([np.abs(np.asarray(np.load(f, allow_pickle=True)["errors_deg"]).ravel()) for f in hfiles_chance])
        b = b[np.isfinite(b)]
        if b.size:
            v2_chance_all.append(b)
            subject_v2_chance_native.append(np.nanmedian(b) if use_median else np.nanmean(b))

# Concatenate all data
all_v1 = np.concatenate(v1_all) if v1_all else np.array([])
all_v1 = all_v1[all_v1 > 0.002]
all_v2 = np.concatenate(v2_all) if v2_all else np.array([])

# Check if chance data is available
has_v1_chance = len(v1_chance_all) > 0
has_v2_chance = len(v2_chance_all) > 0

if has_v1_chance:
    all_v1_chance = np.concatenate(v1_chance_all)
    all_v1_chance = all_v1_chance[all_v1_chance > 0.002]
else:
    all_v1_chance = None
    print("Warning: V1 chance decoding data not found")

if has_v2_chance:
    all_v2_chance = np.concatenate(v2_chance_all)
else:
    all_v2_chance = None
    print("Warning: V2 chance decoding data not found")


# ============ Figure 1: V1 Decoding ============
fig_v1, axs_v1 = plt.subplots(1, 2, figsize=(12, 5))

# Subplot 1: KDE comparison
# Compute plotting range robustly (use combined quantiles so both curves show well)
if all_v1.size:
    if all_v1_chance is not None and all_v1_chance.size:
        _v1_for_range = np.concatenate([all_v1, all_v1_chance])
    else:
        _v1_for_range = all_v1

    qlo, qhi = np.quantile(_v1_for_range[np.isfinite(_v1_for_range)], [0.001, 0.999])
    if not np.isfinite(qlo) or not np.isfinite(qhi) or qhi <= qlo:
        qlo, qhi = float(np.nanmin(_v1_for_range)), float(np.nanmax(_v1_for_range))

    pad = 0.05 * (qhi - qlo) if qhi > qlo else 1.0
    xs_v1 = np.linspace(qlo - pad, qhi + pad, 600)

    # Optional faint histogram fill for context (comment out if you want only KDE)
    axs_v1[0].hist(all_v1, bins=bins, density=density, alpha=0.12, color="C0", edgecolor="none")
    if all_v1_chance is not None:
        axs_v1[0].hist(all_v1_chance, bins=bins, density=density, alpha=0.12, color="C1", edgecolor="none")

    kde_v1 = gaussian_kde(all_v1[np.isfinite(all_v1)])
    axs_v1[0].plot(xs_v1, kde_v1(xs_v1), color="C0", linewidth=2.5, label="Neural Decoder")

    if all_v1_chance is not None and np.isfinite(all_v1_chance).any() and all_v1_chance.size > 1:
        kde_v1_ch = gaussian_kde(all_v1_chance[np.isfinite(all_v1_chance)])
        axs_v1[0].plot(xs_v1, kde_v1_ch(xs_v1), color="C1", linewidth=2.5, linestyle="--",
                       label="Chance (Occupancy)")
else:
    print("Warning: No V1 neural decoding data found")

axs_v1[0].set_xlabel("Absolute V1 Error (cm)")
axs_v1[0].set_ylabel("Density" if density else "Count")
axs_v1[0].set_title("V1 Decoding Error Distribution")
axs_v1[0].legend()

# Subplot 2: Bar plot comparing median errors
v1_median_neural = float(np.nanmedian(all_v1)) if use_median else float(np.nanmean(all_v1))
v1_median_chance = (float(np.nanmedian(all_v1_chance)) if use_median else float(np.nanmean(all_v1_chance))) if all_v1_chance is not None else 0

# Per-subject stats for error bars
subject_v1_neural = np.array(subject_v1_native)
subject_v1_chance = np.array(subject_v1_chance_native)

v1_std_neural = float(np.nanstd(subject_v1_neural, ddof=1)) if len(subject_v1_neural) > 1 else 0
v1_std_chance = float(np.nanstd(subject_v1_chance, ddof=1)) if len(subject_v1_chance) > 1 else 0

bar_labels = ["Neural\nDecoder", "Chance\n(Occupancy)"]
bar_values = [v1_median_neural, v1_median_chance]
bar_errors = [v1_std_neural, v1_std_chance]
bar_colors = ["C0", "C1"]

bars = axs_v1[1].bar([0, 1], bar_values, yerr=bar_errors, capsize=6,
                      alpha=0.8, color=bar_colors)
axs_v1[1].set_xticks([0, 1])
axs_v1[1].set_xticklabels(bar_labels)
axs_v1[1].set_ylabel(f"{'Median' if use_median else 'Mean'} Absolute Error (cm)")
axs_v1[1].set_title("V1 Decoding: Neural vs Chance")

# Add no-prior chance reference line
axs_v1[1].axhline(y=NO_PRIOR_V1_CM, color="C2", linestyle="--", linewidth=1.5,
                  label=f"No Prior ({NO_PRIOR_V1_CM:.1f} cm)")
axs_v1[1].legend(loc="upper right")

# Add value labels on bars
for i, (val, err) in enumerate(zip(bar_values, bar_errors)):
    axs_v1[1].text(i, val + err + 1, f"{val:.2f} cm",
                   ha="center", va="bottom", fontsize=10)

fig_v1.tight_layout()


# ============ Figure 2: V2 Decoding ============
fig_v2, axs_v2 = plt.subplots(1, 2, figsize=(12, 5))

# Subplot 1: KDE comparison
if all_v2.size:
    if all_v2_chance is not None and all_v2_chance.size:
        _v2_for_range = np.concatenate([all_v2, all_v2_chance])
    else:
        _v2_for_range = all_v2

    qlo, qhi = np.quantile(_v2_for_range[np.isfinite(_v2_for_range)], [0.001, 0.999])
    if not np.isfinite(qlo) or not np.isfinite(qhi) or qhi <= qlo:
        qlo, qhi = float(np.nanmin(_v2_for_range)), float(np.nanmax(_v2_for_range))

    pad = 0.05 * (qhi - qlo) if qhi > qlo else 1.0
    xs_v2 = np.linspace(qlo - pad, qhi + pad, 600)

    # Optional faint histogram fill for context (comment out if you want only KDE)
    axs_v2[0].hist(all_v2, bins=bins, density=density, alpha=0.12, color="C0", edgecolor="none")
    if all_v2_chance is not None:
        axs_v2[0].hist(all_v2_chance, bins=bins, density=density, alpha=0.12, color="C1", edgecolor="none")

    kde_v2 = gaussian_kde(all_v2[np.isfinite(all_v2)])
    axs_v2[0].plot(xs_v2, kde_v2(xs_v2), color="C0", linewidth=2.5, label="Neural Decoder")

    if all_v2_chance is not None and np.isfinite(all_v2_chance).any() and all_v2_chance.size > 1:
        kde_v2_ch = gaussian_kde(all_v2_chance[np.isfinite(all_v2_chance)])
        axs_v2[0].plot(xs_v2, kde_v2_ch(xs_v2), color="C1", linewidth=2.5, linestyle="--",
                       label="Chance (Occupancy)")
else:
    print("Warning: No V2 neural decoding data found")

axs_v2[0].set_xlabel("Absolute V2 Error (deg)")
axs_v2[0].set_ylabel("Density" if density else "Count")
axs_v2[0].set_title("V2 Decoding Error Distribution")
axs_v2[0].legend()

# Subplot 2: Bar plot comparing median errors
v2_median_neural = float(np.nanmedian(all_v2)) if use_median else float(np.nanmean(all_v2))
v2_median_chance = (float(np.nanmedian(all_v2_chance)) if use_median else float(np.nanmean(all_v2_chance))) if all_v2_chance is not None else 0

# Per-subject stats for error bars
subject_v2_neural = np.array(subject_v2_native)
subject_v2_chance = np.array(subject_v2_chance_native)

v2_std_neural = float(np.nanstd(subject_v2_neural, ddof=1)) if len(subject_v2_neural) > 1 else 0
v2_std_chance = float(np.nanstd(subject_v2_chance, ddof=1)) if len(subject_v2_chance) > 1 else 0

bar_labels = ["Neural\nDecoder", "Chance\n(Occupancy)"]
bar_values = [v2_median_neural, v2_median_chance]
bar_errors = [v2_std_neural, v2_std_chance]
bar_colors = ["C0", "C1"]

bars = axs_v2[1].bar([0, 1], bar_values, yerr=bar_errors, capsize=6,
                      alpha=0.8, color=bar_colors)
axs_v2[1].set_xticks([0, 1])
axs_v2[1].set_xticklabels(bar_labels)
axs_v2[1].set_ylabel(f"{'Median' if use_median else 'Mean'} Absolute Error (deg)")
axs_v2[1].set_title("V2 Decoding: Neural vs Chance")

# Add no-prior chance reference line
axs_v2[1].axhline(y=NO_PRIOR_V2_DEG, color="C2", linestyle="--", linewidth=1.5,
                  label=f"No Prior ({NO_PRIOR_V2_DEG:.0f} deg)")
axs_v2[1].legend(loc="upper right")

# Add value labels on bars
for i, (val, err) in enumerate(zip(bar_values, bar_errors)):
    axs_v2[1].text(i, val + err + 2, f"{val:.2f} deg",
                   ha="center", va="bottom", fontsize=10)

fig_v2.tight_layout()

# ============ Save figures ============
Path("figs").mkdir(parents=True, exist_ok=True)
fig_v1.savefig("figs/dataset_2_v1_decoding_vs_chance.pdf", dpi=300)
fig_v2.savefig("figs/dataset_2_v2_decoding_vs_chance.pdf", dpi=300)
print("Saved: figs/dataset_2_v1_decoding_vs_chance.pdf")
print("Saved: figs/dataset_2_v2_decoding_vs_chance.pdf")

plt.show()


# CELL 10
save_path = current_file_path / f'results/empirical/dataset_1'
data_path = current_file_path / f'data/empirical/dataset_1'

dataset_n = [1, 2]

# CELL 11
# filter by env type

dataset_1_datasets = {}

from openpyxl import load_workbook

wb = load_workbook(data_path / f"Dataset_metadata.xlsx")

for n in dataset_n:
    sheet = wb[f'Dataset_{n}']

    dataset_dict = {}
    try:
      actual = os.listdir(data_path / f'Dataset_{n}')
    except Exception as e:
      continue
    for row in sheet.iter_rows(values_only=True):
        if not row[0] or row[0][0] != 'A':
            continue
        key = row[0]
        if key not in actual:
            print("Actual dataset missing", key)
        values = (row[3], row[9]) # environment shapes and units
        dataset_dict[key] = values
    dataset_1_datasets[n] = dataset_dict


# CELL 12
from src.dataset_1_decode_solo import decode_solo
import os
import glob
from datetime import datetime, timedelta

reload = True

n_v2_bins = 360
n_x_bins = 32
n_y_bins = 32

def was_processed_recently(save_path, dataset, subject, n_x_bins, n_y_bins, hours=12):
    """Check if v1_response_map files exist and were modified within the last N hours."""
    subject_path = save_path / f'Dataset_{dataset}/{subject}'

    if not subject_path.exists():
        return False

    # Look for files matching the pattern
    pattern = str(subject_path / f'v1_response_map_cv*_{n_x_bins}x{n_y_bins}*')
    matching_files = glob.glob(pattern)

    if not matching_files:
        return False

    # Check if any matching file was modified within the time window
    cutoff = datetime.now() - timedelta(hours=hours)
    for filepath in matching_files:
        mtime = datetime.fromtimestamp(os.path.getmtime(filepath))
        if mtime > cutoff:
            return True

    return False

for n, rats in dataset_1_datasets.items():
    for subject, (env_type, _) in rats.items():
        if env_type == 'Square' or subject == 'A3708-200317b':
            try:
                if not reload and f'errors_v2_cv0_{n_v2_bins}_all.npz' in os.listdir(save_path / f'Dataset_{n}/{subject}'):
                    continue
            except Exception as e:
                pass

            # Skip if already processed recently
            if was_processed_recently(save_path, n, subject, n_x_bins, n_y_bins, hours=12):
                print(f"Skipping Dataset_{n}/{subject} - processed within last 12 hours")
                continue

            decode_solo(
                data_path=data_path,
                save_path=save_path,
                dataset=n,
                subject=subject,
                cv_folds=4,
                save_intermediates=True,
                n_x_bins=32,
                n_y_bins=32
            )

# CELL 14
from tqdm import tqdm

# ============ Configuration ============
root = save_path
bins = 60
density = True
min_v1_cm = 0.5
env_size_cm = 80.0
v1_grid = 32
v2_bins = 360

# No-prior chance values (uniform random guessing)
import math
sqrt2 = math.sqrt(2)
NO_PRIOR_V1_COEFF = (2 + sqrt2 + 5 * math.log(1 + sqrt2)) / 15  # ~0.5214
NO_PRIOR_V1_CM = NO_PRIOR_V1_COEFF * env_size_cm  # ~41.71 cm for 80cm arena
NO_PRIOR_V2_DEG = 90.0  # Expected error for uniform guessing over [0, 360)
NO_PRIOR_V3_DEG = 90.0  # Same for v3

# ============ Imports ============
import numpy as np
import matplotlib.pyplot as plt
from pathlib import Path
from scipy.stats import gaussian_kde

# ============ Load data ============
rats = sorted(
    {p.parent for p in root.rglob("errors_v1_cv*_32x32_all.npz")} |
    {p.parent for p in root.rglob("errors_v2_cv*_360_all.npz")} |
    {p.parent for p in root.rglob("errors_v3_cv*_360_all.npz")}
)

precomp = []
for subject in tqdm(rats):
    # Neural decoding files
    xfiles = list(subject.glob("errors_v1_cv*_32x32_all.npz"))
    hfiles = list(subject.glob("errors_v2_cv*_360_all.npz"))
    mfiles = list(subject.glob("errors_v3_cv*_360_all.npz"))

    # Chance decoding files
    xfiles_chance = list(subject.glob("errors_v1_chance_cv*_32x32_all.npz"))
    hfiles_chance = list(subject.glob("errors_v2_chance_cv*_360_all.npz"))
    mfiles_chance = list(subject.glob("errors_v3_chance_cv*_360_all.npz"))

    if not xfiles or not hfiles:
        continue

    # Load V1 neural decoding errors
    v1 = []
    for f in xfiles:
        d = np.load(f, allow_pickle=True)
        a = np.asarray(d["errors"]).ravel()
        a = a[np.isfinite(a)]
        a = a[a >= min_v1_cm]
        if a.size:
            v1.append(a)

    v1_chance = []
    for f in xfiles_chance:
        d = np.load(f, allow_pickle=True)
        a = np.asarray(d["errors"]).ravel()
        a = a[np.isfinite(a)]
        a = a[a >= min_v1_cm]
        if a.size:
            v1_chance.append(a)

    v2 = []
    for f in hfiles:
        d = np.load(f, allow_pickle=True)
        a = np.asarray(d["errors_deg"]).ravel()
        a = np.abs(a[~np.isnan(a)])
        if a.size:
            v2.append(a)

    v2_chance = []
    for f in hfiles_chance:
        d = np.load(f, allow_pickle=True)
        a = np.asarray(d["errors_deg"]).ravel()
        a = np.abs(a[~np.isnan(a)])
        if a.size:
            v2_chance.append(a)

    # Load V3 neural decoding errors
    v3 = []
    for f in mfiles:
        d = np.load(f, allow_pickle=True)
        a = np.asarray(d["errors_deg"]).ravel()
        a = np.abs(a[~np.isnan(a)])
        if a.size:
            v3.append(a)

    v3_chance = []
    for f in mfiles_chance:
        d = np.load(f, allow_pickle=True)
        a = np.asarray(d["errors_deg"]).ravel()
        a = np.abs(a[~np.isnan(a)])
        if a.size:
            v3_chance.append(a)

    if v1 and v2:
        entry = {
            "v1": np.concatenate(v1),
            "v2": np.concatenate(v2),
        }
        if v1_chance:
            entry["v1_chance"] = np.concatenate(v1_chance)
        if v2_chance:
            entry["v2_chance"] = np.concatenate(v2_chance)
        if v3:
            entry["v3"] = np.concatenate(v3)
        if v3_chance:
            entry["v3_chance"] = np.concatenate(v3_chance)
        precomp.append(entry)

all_v1 = np.concatenate([d["v1"] for d in precomp])
all_v2 = np.concatenate([d["v2"] for d in precomp])

has_v1_chance = all("v1_chance" in d for d in precomp)
has_v2_chance = all("v2_chance" in d for d in precomp)
has_v3 = any("v3" in d for d in precomp)
has_v3_chance = any("v3_chance" in d for d in precomp)

if has_v1_chance:
    all_v1_chance = np.concatenate([d["v1_chance"] for d in precomp])
else:
    all_v1_chance = None
    print("Warning: V1 chance decoding data not found")

if has_v2_chance:
    all_v2_chance = np.concatenate([d["v2_chance"] for d in precomp])
else:
    all_v2_chance = None
    print("Warning: V2 chance decoding data not found")

if has_v3:
    all_v3 = np.concatenate([d["v3"] for d in precomp if "v3" in d])
else:
    all_v3 = None
    print("Warning: V3 decoding data not found")

if has_v3_chance:
    all_v3_chance = np.concatenate([d["v3_chance"] for d in precomp if "v3_chance" in d])
else:
    all_v3_chance = None

# ============ V1 Decoding ============
fig_v1, axs_v1 = plt.subplots(1, 2, figsize=(12, 5))

# Subplot 1: KDE comparison
if all_v1.size:
    if all_v1_chance is not None and all_v1_chance.size:
        _v1_for_range = np.concatenate([all_v1, all_v1_chance])
    else:
        _v1_for_range = all_v1

    _v1_for_range = _v1_for_range[np.isfinite(_v1_for_range)]
    qlo, qhi = np.quantile(_v1_for_range, [0.001, 0.999])
    if not np.isfinite(qlo) or not np.isfinite(qhi) or qhi <= qlo:
        qlo, qhi = float(np.nanmin(_v1_for_range)), float(np.nanmax(_v1_for_range))

    pad = 0.05 * (qhi - qlo) if qhi > qlo else 1.0
    xs_v1 = np.linspace(qlo - pad, qhi + pad, 600)

    # (Optional) faint hist fill for context — remove if you want only KDE lines
    axs_v1[0].hist(all_v1, bins=bins, density=density, alpha=0.12, color="C0", edgecolor="none")
    if all_v1_chance is not None:
        axs_v1[0].hist(all_v1_chance, bins=bins, density=density, alpha=0.12, color="C1", edgecolor="none")

    kde_v1 = gaussian_kde(all_v1[np.isfinite(all_v1)])
    axs_v1[0].plot(xs_v1, kde_v1(xs_v1), color="C0", linewidth=2.5, label="Neural Decoder")

    if all_v1_chance is not None and np.isfinite(all_v1_chance).any() and all_v1_chance.size > 1:
        kde_v1_ch = gaussian_kde(all_v1_chance[np.isfinite(all_v1_chance)])
        axs_v1[0].plot(xs_v1, kde_v1_ch(xs_v1), color="C1", linewidth=2.5, linestyle="--",
                       label="Chance (Occupancy)")

axs_v1[0].set_xlabel("Absolute V1 Error (cm)")
axs_v1[0].set_ylabel("Density" if density else "Count")
axs_v1[0].set_title("V1 Decoding Error Distribution")
axs_v1[0].legend()

v1_median_neural = float(np.nanmedian(all_v1))
v1_median_chance = float(np.nanmedian(all_v1_chance)) if all_v1_chance is not None else 0

subject_v1_neural = np.array([np.nanmedian(d["v1"]) for d in precomp])
subject_v1_chance = np.array([np.nanmedian(d["v1_chance"]) for d in precomp if "v1_chance" in d])

v1_std_neural = float(np.nanstd(subject_v1_neural, ddof=1)) if len(subject_v1_neural) > 1 else 0
v1_std_chance = float(np.nanstd(subject_v1_chance, ddof=1)) if len(subject_v1_chance) > 1 else 0

bar_labels = ["Neural\nDecoder", "Chance\n(Occupancy)"]
bar_values = [v1_median_neural, v1_median_chance]
bar_errors = [v1_std_neural, v1_std_chance]
bar_colors = ["C0", "C1"]

bars = axs_v1[1].bar([0, 1], bar_values, yerr=bar_errors, capsize=6,
                     alpha=0.8, color=bar_colors)
axs_v1[1].set_xticks([0, 1])
axs_v1[1].set_xticklabels(bar_labels)
axs_v1[1].set_ylabel("Median Absolute Error (cm)")
axs_v1[1].set_title("V1 Decoding: Neural vs Chance")

# Add no-prior chance reference line
axs_v1[1].axhline(y=NO_PRIOR_V1_CM, color="C2", linestyle="--", linewidth=1.5,
                  label=f"No Prior ({NO_PRIOR_V1_CM:.1f} cm)")
axs_v1[1].legend(loc="upper right")

# Add value labels on bars
for i, (val, err) in enumerate(zip(bar_values, bar_errors)):
    axs_v1[1].text(i, val + err + 0.5, f"{val:.2f} cm",
                   ha="center", va="bottom", fontsize=10)
fig_v1.tight_layout()

# ============ V2 Decoding ============
fig_v2, axs_v2 = plt.subplots(1, 2, figsize=(12, 5))

# Subplot 1: KDE comparison
if all_v2.size:
    if all_v2_chance is not None and all_v2_chance.size:
        _v2_for_range = np.concatenate([all_v2, all_v2_chance])
    else:
        _v2_for_range = all_v2

    _v2_for_range = _v2_for_range[np.isfinite(_v2_for_range)]
    qlo, qhi = np.quantile(_v2_for_range, [0.001, 0.999])
    if not np.isfinite(qlo) or not np.isfinite(qhi) or qhi <= qlo:
        qlo, qhi = float(np.nanmin(_v2_for_range)), float(np.nanmax(_v2_for_range))

    pad = 0.05 * (qhi - qlo) if qhi > qlo else 1.0
    xs_v2 = np.linspace(qlo - pad, qhi + pad, 600)

    # (Optional) faint hist fill for context — remove if you want only KDE lines
    axs_v2[0].hist(all_v2, bins=bins, density=density, alpha=0.12, color="C0", edgecolor="none")
    if all_v2_chance is not None:
        axs_v2[0].hist(all_v2_chance, bins=bins, density=density, alpha=0.12, color="C1", edgecolor="none")

    kde_v2 = gaussian_kde(all_v2[np.isfinite(all_v2)])
    axs_v2[0].plot(xs_v2, kde_v2(xs_v2), color="C0", linewidth=2.5, label="Neural Decoder")

    if all_v2_chance is not None and np.isfinite(all_v2_chance).any() and all_v2_chance.size > 1:
        kde_v2_ch = gaussian_kde(all_v2_chance[np.isfinite(all_v2_chance)])
        axs_v2[0].plot(xs_v2, kde_v2_ch(xs_v2), color="C1", linewidth=2.5, linestyle="--",
                       label="Chance (Occupancy)")

axs_v2[0].set_xlabel("Absolute V2 Error (deg)")
axs_v2[0].set_ylabel("Density" if density else "Count")
axs_v2[0].set_title("V2 Decoding Error Distribution")
axs_v2[0].legend()

v2_median_neural = float(np.nanmedian(all_v2))
v2_median_chance = float(np.nanmedian(all_v2_chance)) if all_v2_chance is not None else 0

subject_v2_neural = np.array([np.nanmedian(d["v2"]) for d in precomp])
subject_v2_chance = np.array([np.nanmedian(d["v2_chance"]) for d in precomp if "v2_chance" in d])

v2_std_neural = float(np.nanstd(subject_v2_neural, ddof=1)) if len(subject_v2_neural) > 1 else 0
v2_std_chance = float(np.nanstd(subject_v2_chance, ddof=1)) if len(subject_v2_chance) > 1 else 0

bar_labels = ["Neural\nDecoder", "Chance\n(Occupancy)"]
bar_values = [v2_median_neural, v2_median_chance]
bar_errors = [v2_std_neural, v2_std_chance]
bar_colors = ["C0", "C1"]

bars = axs_v2[1].bar([0, 1], bar_values, yerr=bar_errors, capsize=6,
                     alpha=0.8, color=bar_colors)
axs_v2[1].set_xticks([0, 1])
axs_v2[1].set_xticklabels(bar_labels)
axs_v2[1].set_ylabel("Median Absolute Error (deg)")
axs_v2[1].set_title("V2 Decoding: Neural vs Chance")

# Add no-prior chance reference line
axs_v2[1].axhline(y=NO_PRIOR_V2_DEG, color="C2", linestyle="--", linewidth=1.5,
                  label=f"No Prior ({NO_PRIOR_V2_DEG:.0f} deg)")
axs_v2[1].legend(loc="upper right")

for i, (val, err) in enumerate(zip(bar_values, bar_errors)):
    axs_v2[1].text(i, val + err + 1, f"{val:.2f} deg",
                   ha="center", va="bottom", fontsize=10)

fig_v2.tight_layout()

# ============ V3 Decoding ============
if all_v3 is not None:
    fig_v3, axs_v3 = plt.subplots(1, 2, figsize=(12, 5))

    # Subplot 1: KDE comparison
    if all_v3.size:
        if all_v3_chance is not None and all_v3_chance.size:
            _v3_for_range = np.concatenate([all_v3, all_v3_chance])
        else:
            _v3_for_range = all_v3

        _v3_for_range = _v3_for_range[np.isfinite(_v3_for_range)]
        qlo, qhi = np.quantile(_v3_for_range, [0.001, 0.999])
        if not np.isfinite(qlo) or not np.isfinite(qhi) or qhi <= qlo:
            qlo, qhi = float(np.nanmin(_v3_for_range)), float(np.nanmax(_v3_for_range))

        pad = 0.05 * (qhi - qlo) if qhi > qlo else 1.0
        xs_v3 = np.linspace(qlo - pad, qhi + pad, 600)

        # (Optional) faint hist fill for context — remove if you want only KDE lines
        axs_v3[0].hist(all_v3, bins=bins, density=density, alpha=0.12, color="C0", edgecolor="none")
        if all_v3_chance is not None:
            axs_v3[0].hist(all_v3_chance, bins=bins, density=density, alpha=0.12, color="C1", edgecolor="none")

        kde_v3 = gaussian_kde(all_v3[np.isfinite(all_v3)])
        axs_v3[0].plot(xs_v3, kde_v3(xs_v3), color="C0", linewidth=2.5, label="Neural Decoder")

        if all_v3_chance is not None and np.isfinite(all_v3_chance).any() and all_v3_chance.size > 1:
            kde_v3_ch = gaussian_kde(all_v3_chance[np.isfinite(all_v3_chance)])
            axs_v3[0].plot(xs_v3, kde_v3_ch(xs_v3), color="C1", linewidth=2.5, linestyle="--",
                           label="Chance (Occupancy)")

    axs_v3[0].set_xlabel("Absolute V3 Error (deg)")
    axs_v3[0].set_ylabel("Density" if density else "Count")
    axs_v3[0].set_title("V3 Decoding Error Distribution")
    axs_v3[0].legend()

    v3_median_neural = float(np.nanmedian(all_v3))
    v3_median_chance = float(np.nanmedian(all_v3_chance)) if all_v3_chance is not None else 0

    subject_v3_neural = np.array([np.nanmedian(d["v3"]) for d in precomp if "v3" in d])
    subject_v3_chance = np.array([np.nanmedian(d["v3_chance"]) for d in precomp if "v3_chance" in d])

    v3_std_neural = float(np.nanstd(subject_v3_neural, ddof=1)) if len(subject_v3_neural) > 1 else 0
    v3_std_chance = float(np.nanstd(subject_v3_chance, ddof=1)) if len(subject_v3_chance) > 1 else 0

    bar_labels = ["Neural\nDecoder", "Chance\n(Occupancy)"]
    bar_values = [v3_median_neural, v3_median_chance]
    bar_errors = [v3_std_neural, v3_std_chance]
    bar_colors = ["C0", "C1"]

    bars = axs_v3[1].bar([0, 1], bar_values, yerr=bar_errors, capsize=6,
                         alpha=0.8, color=bar_colors)
    axs_v3[1].set_xticks([0, 1])
    axs_v3[1].set_xticklabels(bar_labels)
    axs_v3[1].set_ylabel("Median Absolute Error (deg)")
    axs_v3[1].set_title("V3 Decoding: Neural vs Chance")

    # Add no-prior chance reference line
    axs_v3[1].axhline(y=NO_PRIOR_V3_DEG, color="C2", linestyle="--", linewidth=1.5,
                      label=f"No Prior ({NO_PRIOR_V3_DEG:.0f} deg)")
    axs_v3[1].legend(loc="upper right")

    for i, (val, err) in enumerate(zip(bar_values, bar_errors)):
        axs_v3[1].text(i, val + err + 1, f"{val:.2f} deg",
                       ha="center", va="bottom", fontsize=10)

    fig_v3.tight_layout()

# ============ Save figures ============
Path("figs").mkdir(parents=True, exist_ok=True)
fig_v1.savefig("figs/dataset_1_v1_decoding_vs_chance.pdf", dpi=300)
fig_v2.savefig("figs/dataset_1_v2_decoding_vs_chance.pdf", dpi=300)
print("Saved: figs/dataset_1_v1_decoding_vs_chance.pdf")
print("Saved: figs/dataset_1_v2_decoding_vs_chance.pdf")

if all_v3 is not None:
    fig_v3.savefig("figs/dataset_1_v3_decoding_vs_chance.pdf", dpi=300)
    print("Saved: figs/dataset_1_v3_decoding_vs_chance.pdf")

plt.show()


# CELL 16
from pathlib import Path
from typing import Dict, List, Optional, Literal
import numpy as np
import matplotlib.pyplot as plt
from matplotlib.backends.backend_pdf import PdfPages
import math
from scipy.stats import gaussian_kde

# ============Configuration ============
dataset = "dataset_1"  # "dataset_2" or "dataset_1"
mode = "v2_v3"    # "v1_v2" or "v2_v3" (v2_v3 only for dataset_1)
# save_path = Path("./results/dataset_2")  # for dataset_2
# save_path already defined in notebook for dataset_1

use_median = True
bins = 60
min_units = 40

# No-prior chance values (uniform random guessing)
sqrt2 = math.sqrt(2)
NO_PRIOR_V1_COEFF = (2 + sqrt2 + 5 * math.log(1 + sqrt2)) / 15  # ~0.5214
# Environment sizes
DATASET_2_ENV_CM = 200.0
DATASET_1_ENV_CM = 80.0
NO_PRIOR_V1_DATASET_2 = NO_PRIOR_V1_COEFF * DATASET_2_ENV_CM  # ~104.28 cm
NO_PRIOR_V1_DATASET_1 = NO_PRIOR_V1_COEFF * DATASET_1_ENV_CM  # ~41.71 cm
NO_PRIOR_V2_DEG = 90.0  # Expected error for uniform guessing over [0, 360)
NO_PRIOR_V3_DEG = 90.0  # Same for v3

# ============ Helpers ============

def plot_error_comparison(ax_hist, ax_bar, errors_neural, errors_chance, label, unit, no_prior_val=None):
    # ---- KDE plot (replaces histogram bars) ----
    errors_neural = np.asarray(errors_neural)
    errors_neural = errors_neural[np.isfinite(errors_neural)]
    errors_ch = None
    if errors_chance is not None and len(errors_chance) > 0:
        errors_ch = np.asarray(errors_chance)
        errors_ch = errors_ch[np.isfinite(errors_ch)]

    if errors_ch is not None and errors_ch.size:
        _for_range = np.concatenate([errors_neural, errors_ch])
    else:
        _for_range = errors_neural

    if _for_range.size:
        qlo, qhi = np.quantile(_for_range, [0.001, 0.999])
        if not np.isfinite(qlo) or not np.isfinite(qhi) or qhi <= qlo:
            qlo, qhi = float(np.nanmin(_for_range)), float(np.nanmax(_for_range))

        pad = 0.05 * (qhi - qlo) if qhi > qlo else 1.0
        xs = np.linspace(qlo - pad, qhi + pad, 600)

        # (Optional) faint hist fill for context — delete these lines if you want KDE-only
        ax_hist.hist(errors_neural, bins=bins, density=True, alpha=0.12, color="C0", edgecolor="none")
        if errors_ch is not None:
            ax_hist.hist(errors_ch, bins=bins, density=True, alpha=0.12, color="C1", edgecolor="none")

        kde_n = gaussian_kde(errors_neural)
        ax_hist.plot(xs, kde_n(xs), color="C0", linewidth=2.5, label="Neural Decoder")

        if errors_ch is not None and errors_ch.size > 1:
            kde_c = gaussian_kde(errors_ch)
            ax_hist.plot(xs, kde_c(xs), color="C1", linewidth=2.5, linestyle="--", label="Chance (Occupancy)")

    ax_hist.set_xlabel(f"Absolute {label} Error ({unit})")
    ax_hist.set_ylabel("Density")
    ax_hist.set_title(f"{label} Error Distribution")
    ax_hist.legend()

    # ---- bar plot (unchanged logic, but handle None safely) ----
    stat_fn = np.nanmedian if use_median else np.nanmean
    neural_val = float(stat_fn(errors_neural)) if errors_neural.size else 0.0
    chance_val = float(stat_fn(errors_ch)) if (errors_ch is not None and errors_ch.size) else 0.0

    ax_bar.bar([0, 1], [neural_val, chance_val], alpha=0.8, color=["C0", "C1"])
    ax_bar.set_xticks([0, 1])
    ax_bar.set_xticklabels(["Neural", "Chance"])
    ax_bar.set_ylabel(f"{'Median' if use_median else 'Mean'} Error ({unit})")
    ax_bar.set_title(f"{label}: Neural vs Chance")

    # Add no-prior chance reference line if provided
    if no_prior_val is not None:
        ax_bar.axhline(y=no_prior_val, color="C2", linestyle="--", linewidth=1.5,
                       label=f"No Prior ({no_prior_val:.1f})")
        ax_bar.legend(loc="upper right")

    # Value labels
    denom = max(neural_val, chance_val, 1e-9)
    for i, val in enumerate([neural_val, chance_val]):
        ax_bar.text(i, val + denom * 0.05, f"{val:.2f}{unit}", ha="center", va="bottom")


def load_errors(files, key="errors_deg", take_abs=True):
    arrays = []
    for f in files:
        try:
            arr = np.asarray(np.load(f, allow_pickle=True)[key]).ravel()
            arr = arr[np.isfinite(arr)]
            if take_abs:
                arr = np.abs(arr)
            if key == 'errors' or key == 'errors_cm':
                arr = arr[arr > 0.2]
            if arr.size > 0:
                arrays.append(arr)
        except:
            pass
    return np.concatenate(arrays) if arrays else None


# ============ Discover sessions ============

sessions = []

if dataset == "dataset_2":
    subject_dirs = sorted({f.parent for f in save_path.rglob("errors_v1_cv*_*x*.npz")} |
                      {f.parent for f in save_path.rglob("errors_v2_cv*_*.npz")})

    for dpath in subject_dirs:
        name = dpath.relative_to(save_path).as_posix()
        v1_neural = load_errors([f for f in dpath.glob("errors_v1_cv*_*x*.npz") if "chance" not in f.name], key="errors_cm")
        v1_chance = load_errors(list(dpath.glob("errors_v1_chance_cv*_*x*.npz")), key="errors_cm")
        v2_neural = load_errors([f for f in dpath.glob("errors_v2_cv*_*.npz") if "chance" not in f.name], key="errors_deg")
        v2_chance = load_errors(list(dpath.glob("errors_v2_chance_cv*_*.npz")), key="errors_deg")

        if v1_neural is not None or v2_neural is not None:
            sessions.append({
                'name': name, 'var1': v1_neural, 'var1_c': v1_chance, 'var2': v2_neural, 'var2_c': v2_chance,
                'l1': 'V1', 'u1': 'cm', 'l2': 'V2', 'u2': 'deg',
                'no_prior_1': NO_PRIOR_V1_DATASET_2, 'no_prior_2': NO_PRIOR_V2_DEG
            })

else:  # dataset_1
    if mode == "v1_v2":
        subject_dirs = sorted({f.parent for f in save_path.rglob("errors_v1_cv*_*x*_all.npz")} |
                          {f.parent for f in save_path.rglob("errors_v2_cv*_*_all.npz")})
    else:
        subject_dirs = sorted({f.parent for f in save_path.rglob("errors_v2_cv*_*_all.npz")} |
                          {f.parent for f in save_path.rglob("errors_v3_cv*_*_all.npz")})

    for dpath in subject_dirs:
        name = dpath.relative_to(save_path).as_posix()

        if min_units > 0:
            kept_path = dpath / "units_kept_v2_filter.npy"
            if kept_path.exists():
                try:
                    n = float(np.unique(np.load(kept_path, allow_pickle=True)).size)
                    if n < min_units:
                        continue
                except:
                    pass

        if mode == "v1_v2":
            var1 = load_errors([f for f in dpath.glob("errors_v1_cv*_*x*_all.npz") if "chance" not in f.name], key="errors")
            var1_c = load_errors(list(dpath.glob("errors_v1_chance_cv*_*x*_all.npz")), key="errors")
            var2 = load_errors([f for f in dpath.glob("errors_v2_cv*_*_all.npz") if "chance" not in f.name], key="errors_deg")
            var2_c = load_errors(list(dpath.glob("errors_v2_chance_cv*_*_all.npz")), key="errors_deg")
            l1, u1, l2, u2 = "V1", "cm", "V2", "deg"
            no_prior_1, no_prior_2 = NO_PRIOR_V1_DATASET_1, NO_PRIOR_V2_DEG
        else:
            var1 = load_errors([f for f in dpath.glob("errors_v2_cv*_*_all.npz") if "chance" not in f.name], key="errors_deg")
            var1_c = load_errors(list(dpath.glob("errors_v2_chance_cv*_*_all.npz")), key="errors_deg")
            var2 = load_errors([f for f in dpath.glob("errors_v3_cv*_*_all.npz") if "chance" not in f.name], key="errors_deg")
            var2_c = load_errors(list(dpath.glob("errors_v3_chance_cv*_*_all.npz")), key="errors_deg")
            l1, u1, l2, u2 = "V2", "deg", "V3", "deg"
            no_prior_1, no_prior_2 = NO_PRIOR_V2_DEG, NO_PRIOR_V3_DEG

        if var1 is not None or var2 is not None:
            sessions.append({'name': name, 'var1': var1, 'var1_c': var1_c, 'var2': var2, 'var2_c': var2_c,
                             'l1': l1, 'u1': u1, 'l2': l2, 'u2': u2,
                             'no_prior_1': no_prior_1, 'no_prior_2': no_prior_2})

print(f"Found {len(sessions)} sessions")

# ============ Generate PDF ============

output_dir = save_path / "figs"
output_dir.mkdir(parents=True, exist_ok=True)
output_path = output_dir / f"{dataset}_{mode}_individual_sessions.pdf"

with PdfPages(output_path) as pdf:
    for i, s in enumerate(sessions):
        print(f"  [{i+1}/{len(sessions)}] {s['name']}")

        if s['var1'] is None and s['var2'] is None:
            continue

        fig, axs = plt.subplots(2, 2, figsize=(12, 10))
        fig.suptitle(s['name'], fontsize=14, fontweight='bold')

        if s['var1'] is not None and len(s['var1']) > 0:
            plot_error_comparison(axs[0, 0], axs[0, 1], s['var1'], s.get('var1_c'), s['l1'], s['u1'], s.get('no_prior_1'))
        if s['var2'] is not None and len(s['var2']) > 0:
            plot_error_comparison(axs[1, 0], axs[1, 1], s['var2'], s.get('var2_c'), s['l2'], s['u2'], s.get('no_prior_2'))

        fig.tight_layout(rect=[0, 0, 1, 0.96])
        pdf.savefig(fig, dpi=150)
        plt.close(fig)

print(f"\nSaved: {output_path}")


# CELL 18
# %% [markdown]
# Collect a dictionary of session titles -> V2 unit counts for Dataset_1 dataset

# %%
from pathlib import Path
import numpy as np
from scipy.io import loadmat

# -- robust MATLAB struct extractor (matches your pipeline) --
def _extract_matlab_obj(obj):
    try:
        from scipy.io.matlab._mio5_params import mat_struct
    except Exception:
        mat_struct = ()  # fallback if internal path changes
    if not (isinstance(obj, dict) or isinstance(obj, mat_struct)):
        return obj
    d = obj.__dict__ if hasattr(obj, "__dict__") else obj
    out = {}
    for k, v in d.items():
        if k in ["_fieldnames", "type", "time_unit", "data_unit", "name"]:
            continue
        if hasattr(v, "__dict__") or isinstance(v, dict):
            out[k] = _extract_matlab_obj(v)
        elif isinstance(v, np.ndarray):
            out[k] = np.array([_extract_matlab_obj(x) for x in v])
        else:
            out[k] = v
    return out

def _to_1d_int(arr):
    a = np.asarray(arr).ravel()
    # convert booleans or truthy to ints
    if a.dtype == bool:
        a = a.astype(int)
    return a.astype(int, copy=False)

def count_v2_units(celltypes_mat_path: Path) -> int:
    """Return number of V2 units (v2 == 1) from a Dataset_1 CellTypes.mat file."""
    ct_raw = loadmat(celltypes_mat_path, struct_as_record=False, squeeze_me=True)
    # Commonly the struct is under 'CellTypes', but support variants.
    ct = _extract_matlab_obj(ct_raw.get('CellTypes', ct_raw))
    # Try several likely places for the v2 vector
    candidates = [
        ct.get('v2'),
        ct_raw.get('v2'),
        (ct.get('labels', {}) if isinstance(ct.get('labels'), dict) else {}).get('v2'),
    ]
    v2_vec = None
    for c in candidates:
        if c is not None and np.size(c) > 0:
            v2_vec = _to_1d_int(c)
            break
    if v2_vec is None:
        raise FileNotFoundError(f"Couldn't find 'v2' labels in {celltypes_mat_path}")
    # Convention: v2 == 1 means V2 unit
    return int(np.sum(v2_vec == 1))

def gather_v2_counts(
    data_root: Path = Path("./data/dataset_1"),
    sessions: list[tuple[int, str]] | None = None,
) -> dict[str, int]:
    """
    Build { 'Dataset_{id}/SUBJECT': v2_count }.
    If `sessions` is None, auto-discovers all Dataset_*/*/Data/CellTypes.mat.
    """
    results: dict[str, int] = {}
    if sessions is None:
        for ct_path in sorted(data_root.glob("Dataset_*/*/Data/CellTypes.mat")):
            dataset_str = ct_path.parts[-4]  # "Dataset_{id}"
            subject = ct_path.parts[-3]
            title = f"{dataset_str}/{subject}"
            try:
                results[title] = count_v2_units(ct_path)
            except Exception as e:
                results[title] = -1  # mark failure
    else:
        for ds_id, subject in sessions:
            ct_path = data_root / f"Dataset_{ds_id}" / subject / "Data" / "CellTypes.mat"
            title = f"Dataset_{ds_id}/{subject}"
            try:
                results[title] = count_v2_units(ct_path)
            except Exception:
                results[title] = -1
    return results


v2_counts = gather_v2_counts(Path(data_path), sessions=session_list)

# Print the dictionary (and a neat summary)
print(v2_counts)
print("\nSummary:")
ok = {k: v for k, v in v2_counts.items() if v >= 0}
fail = [k for k, v in v2_counts.items() if v < 0]
print(f"  Sessions found: {len(v2_counts)}")
print(f"  With counts    : {len(ok)} (total V2 units = {sum(ok.values())})")
if fail:
    print(f"  Failed to read : {len(fail)} -> {fail}")


import json
from pathlib import Path

# --- after computing v2_counts ---
output_path = Path(save_path / "v2_unit_counts.json")

# Save dictionary to JSON
with open(output_path, "w") as f:
    json.dump(v2_counts, f, indent=2)  # indent makes it nicely formatted

print(f"Saved V2 counts to {output_path}")

# --- later on, when you want to load it ---
with open(output_path, "r") as f:
    loaded_v2_counts = json.load(f)

print("Loaded V2 counts:", loaded_v2_counts)



# CELL 19
